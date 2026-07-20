from pathlib import Path

import pytest

from rag_agent_platform.ingestion.loaders import (
    DocxDocumentLoader,
    LoaderRegistry,
    PdfDocumentLoader,
    TextDocumentLoader,
    build_default_loader_registry,
)


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("leave_policy.txt", "Leave policy\nAnnual leave requests need approval."),
        ("purchase_process.md", "# Purchase process\nDepartment owner approves purchase requests."),
    ],
)
def test_text_and_markdown_can_be_loaded(
    tmp_path: Path,
    filename: str,
    content: str,
) -> None:
    file_path = tmp_path / filename
    file_path.write_text(content, encoding="utf-8")

    loaded = build_default_loader_registry().load(file_path)

    assert loaded.content == content
    assert loaded.metadata["filename"] == filename
    assert loaded.metadata["file_type"] == file_path.suffix.lower().lstrip(".")
    assert loaded.metadata["encoding"] == "utf-8"
    assert loaded.metadata["source_path"].endswith(filename)


def test_gbk_text_can_be_loaded(tmp_path: Path) -> None:
    file_path = tmp_path / "gbk-policy.txt"
    content = "\u8d39\u7528\u62a5\u9500\u5236\u5ea6"
    file_path.write_text(content, encoding="gbk")

    loaded = build_default_loader_registry().load(file_path)

    assert loaded.content == content
    assert loaded.metadata["encoding"] in {"gbk", "gb18030"}


def test_pdf_can_be_loaded(tmp_path: Path) -> None:
    file_path = tmp_path / "sample.pdf"
    write_simple_pdf(file_path, "PDF policy requires review.")

    loaded = build_default_loader_registry().load(file_path)

    assert "PDF policy requires review." in loaded.content
    assert loaded.metadata["filename"] == "sample.pdf"
    assert loaded.metadata["file_type"] == "pdf"
    assert loaded.metadata["page_count"] == 1


def test_docx_can_be_loaded(tmp_path: Path) -> None:
    from docx import Document

    file_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("DOCX policy requires approval.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "HR"
    document.save(file_path)

    loaded = build_default_loader_registry().load(file_path)

    assert "DOCX policy requires approval." in loaded.content
    assert "Owner" in loaded.content
    assert "HR" in loaded.content
    assert loaded.metadata["filename"] == "sample.docx"
    assert loaded.metadata["file_type"] == "docx"
    assert loaded.metadata["paragraph_count"] >= 1
    assert loaded.metadata["table_count"] == 1


def test_empty_file_fails(tmp_path: Path) -> None:
    file_path = tmp_path / "empty.txt"
    file_path.write_text("  \n", encoding="utf-8")

    with pytest.raises(ValueError, match="must not be empty"):
        build_default_loader_registry().load(file_path)


def test_missing_file_fails(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="does not exist"):
        build_default_loader_registry().load(tmp_path / "missing.txt")


def test_unsupported_file_type_fails(tmp_path: Path) -> None:
    file_path = tmp_path / "image.png"
    file_path.write_bytes(b"not a document")

    with pytest.raises(ValueError, match="unsupported file type"):
        build_default_loader_registry().load(file_path)


def test_loader_registry_routes_registered_suffix(tmp_path: Path) -> None:
    file_path = tmp_path / "sample.txt"
    file_path.write_text("sample text", encoding="utf-8")
    registry = LoaderRegistry([TextDocumentLoader()])

    loaded = registry.load(file_path)

    assert loaded.metadata["filename"] == "sample.txt"


def test_default_registry_supports_text_pdf_and_docx() -> None:
    registry = build_default_loader_registry()

    assert isinstance(registry._loaders[".txt"], TextDocumentLoader)
    assert isinstance(registry._loaders[".md"], TextDocumentLoader)
    assert isinstance(registry._loaders[".pdf"], PdfDocumentLoader)
    assert isinstance(registry._loaders[".docx"], DocxDocumentLoader)


def write_simple_pdf(path: Path, text: str) -> None:
    escaped_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped_text}) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream.encode())} >>\nstream\n{stream}\nendstream".encode(),
    ]

    parts = [b"%PDF-1.4\n"]
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(sum(len(part) for part in parts))
        parts.append(f"{index} 0 obj\n".encode())
        parts.append(obj)
        parts.append(b"\nendobj\n")

    xref_offset = sum(len(part) for part in parts)
    parts.append(f"xref\n0 {len(objects) + 1}\n".encode())
    parts.append(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        parts.append(f"{offset:010d} 00000 n \n".encode())
    parts.append(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    path.write_bytes(b"".join(parts))
